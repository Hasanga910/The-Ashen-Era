from pathlib import Path
import json
import hashlib
import re

import fitz
from docx import Document

from indexing.chunker import chunk_page_text


# ============================================================
# PATHS
# ============================================================

ROOT = Path("data/raw/Ashen_Era_Archive")
PROCESSED = Path("data/processed")
VISUALS = Path("data/visuals")

PROCESSED.mkdir(parents=True, exist_ok=True)
VISUALS.mkdir(parents=True, exist_ok=True)


# ============================================================
# HELPERS
# ============================================================

def safe_name(path):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", path)


def make_visual_id(source, page, index):
    raw = f"{source}|{page}|{index}".encode("utf-8")
    return hashlib.md5(raw).hexdigest()[:12]


def make_standalone_visual_id(relative_path):
    return hashlib.md5(
        relative_path.encode("utf-8")
    ).hexdigest()[:12]


def resolve_markdown_image(path, image_ref):
    """
    Resolve an image referenced inside a Markdown file.

    First try:
        Markdown file's own folder / image_ref

    Then try:
        ROOT / image_ref
    """

    image_ref = image_ref.strip()

    # Remove optional surrounding angle brackets
    if image_ref.startswith("<") and image_ref.endswith(">"):
        image_ref = image_ref[1:-1]

    possible_paths = [
        path.parent / image_ref,
        ROOT / image_ref
    ]

    for image_path in possible_paths:
        try:
            image_path = image_path.resolve()

            if image_path.exists() and image_path.is_file():
                return image_path

        except Exception:
            continue

    return None


# ============================================================
# PDF
# ============================================================

def process_pdf(path):
    print(f"[PDF] {path}")

    doc = fitz.open(path)
    all_chunks = []

    for page_index, page in enumerate(doc):

        page_number = page_index + 1
        text = page.get_text("text").strip()

        image_ids = []

        # ----------------------------------------------------
        # Extract embedded images
        # ----------------------------------------------------

        try:
            image_infos = page.get_image_info(xrefs=True)

            for image_index, info in enumerate(image_infos):

                bbox = info.get("bbox")

                if not bbox:
                    continue

                try:
                    rect = fitz.Rect(bbox)

                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(2, 2),
                        clip=rect,
                        alpha=False
                    )

                    visual_id = make_visual_id(
                        str(path.relative_to(ROOT)),
                        page_number,
                        image_index
                    )

                    filename = (
                        f"{safe_name(path.stem)}"
                        f"_p{page_number}"
                        f"_img{image_index}"
                        f"_{visual_id}.png"
                    )

                    output = VISUALS / filename

                    pix.save(output)

                    image_ids.append(visual_id)

                except Exception as e:
                    print(
                        f"    [WARN] Image extraction failed: {e}"
                    )

        except Exception as e:
            print(
                f"    [WARN] Could not inspect images: {e}"
            )

        # ----------------------------------------------------
        # Text chunks
        # ----------------------------------------------------

        if text:

            chunks = chunk_page_text(
                str(path.relative_to(ROOT)),
                page_number,
                text
            )

            # Attach first visual on this page
            # to the first text chunk.
            if image_ids and chunks:
                chunks[0].linked_image_id = image_ids[0]

            all_chunks.extend(chunks)

    # --------------------------------------------------------
    # Save PDF chunks
    # --------------------------------------------------------

    output_name = safe_name(path.stem) + "_chunks.json"
    output_path = PROCESSED / output_name

    with open(output_path, "w", encoding="utf-8") as f:

        json.dump(
            [c.__dict__ for c in all_chunks],
            f,
            ensure_ascii=False,
            indent=2
        )

    print(f"    Pages: {len(doc)}")
    print(f"    Chunks: {len(all_chunks)}")
    print(f"    Saved: {output_path}")

    return len(all_chunks)


# ============================================================
# MARKDOWN
# ============================================================

def process_markdown(path):
    print(f"[MD]  {path}")

    text = path.read_text(
        encoding="utf-8",
        errors="replace"
    ).strip()

    if not text:
        return 0

    # --------------------------------------------------------
    # Find ALL Markdown image references
    #
    # Example:
    # ![House Morvain](images/house_morvain.png)
    # --------------------------------------------------------

    image_matches = re.findall(
        r'!\[[^\]]*\]\(([^)\s]+)(?:\s+["\'][^"\']*["\'])?\)',
        text
    )

    resolved_images = []

    for image_ref in image_matches:

        image_path = resolve_markdown_image(
            path,
            image_ref
        )

        if image_path is None:
            print(
                f"    [WARN] Image not found: {image_ref}"
            )
            continue

        try:
            relative_image_path = str(
                image_path.relative_to(ROOT)
            )
        except ValueError:
            print(
                f"    [WARN] Image outside corpus: {image_path}"
            )
            continue

        visual_id = make_standalone_visual_id(
            relative_image_path
        )

        resolved_images.append(
            {
                "image_ref": image_ref,
                "image_path": image_path,
                "relative_path": relative_image_path,
                "visual_id": visual_id
            }
        )

    # --------------------------------------------------------
    # Remove Markdown image syntax from text.
    #
    # We do NOT want Streamlit to try rendering:
    #
    # ![House Morvain](images/...)
    #
    # Instead, Streamlit will use linked_image_id.
    # --------------------------------------------------------

    clean_text = re.sub(
        r'!\[[^\]]*\]\(([^)\s]+)(?:\s+["\'][^"\']*["\'])?\)',
        "",
        text
    )

    clean_text = re.sub(
        r'\n{3,}',
        '\n\n',
        clean_text
    ).strip()

    # --------------------------------------------------------
    # Create normal text chunks
    # --------------------------------------------------------

    chunks = []

    if clean_text:

        text_chunks = chunk_page_text(
            str(path.relative_to(ROOT)),
            None,
            clean_text
        )

        chunks.extend(text_chunks)

    # --------------------------------------------------------
    # Create a dedicated chunk for EVERY image.
    #
    # This is the important part.
    # --------------------------------------------------------

    for index, image in enumerate(resolved_images):

        image_chunk = {
            "chunk_id": (
                f"image_"
                f"{image['visual_id']}_"
                f"{index}"
            ),

            "source_file": str(
                path.relative_to(ROOT)
            ),

            "page_number": None,

            "text": (
                f"Image asset associated with "
                f"{path.name}: "
                f"{image['relative_path']}"
            ),

            "linked_image_id": image["visual_id"]
        }

        chunks.append(
            type(
                "ImageChunk",
                (),
                image_chunk
            )()
        )

    # --------------------------------------------------------
    # Save Markdown chunks
    # --------------------------------------------------------

    output_name = safe_name(path.stem) + "_chunks.json"
    output_path = PROCESSED / output_name

    with open(output_path, "w", encoding="utf-8") as f:

        json.dump(
            [
                c.__dict__
                if hasattr(c, "__dict__")
                else c
                for c in chunks
            ],
            f,
            ensure_ascii=False,
            indent=2
        )

    print(f"    Text chunks: {len(chunks) - len(resolved_images)}")
    print(f"    Linked visuals: {len(resolved_images)}")
    print(f"    Total chunks: {len(chunks)}")
    print(f"    Saved: {output_path}")

    return len(chunks)


# ============================================================
# TXT
# ============================================================

def process_txt(path):
    print(f"[TXT] {path}")

    text = path.read_text(
        encoding="utf-8",
        errors="replace"
    ).strip()

    if not text:
        return 0

    chunks = chunk_page_text(
        str(path.relative_to(ROOT)),
        None,
        text
    )

    output_name = safe_name(path.stem) + "_chunks.json"
    output_path = PROCESSED / output_name

    with open(output_path, "w", encoding="utf-8") as f:

        json.dump(
            [c.__dict__ for c in chunks],
            f,
            ensure_ascii=False,
            indent=2
        )

    print(f"    Chunks: {len(chunks)}")
    print(f"    Saved: {output_path}")

    return len(chunks)


# ============================================================
# DOCX
# ============================================================

def process_docx(path):
    print(f"[DOCX] {path}")

    doc = Document(path)

    parts = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in doc.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(cells)

            if row_text.strip():
                parts.append(row_text)

    text = "\n\n".join(parts).strip()

    if not text:
        return 0

    chunks = chunk_page_text(
        str(path.relative_to(ROOT)),
        None,
        text
    )

    output_name = safe_name(path.stem) + "_chunks.json"
    output_path = PROCESSED / output_name

    with open(output_path, "w", encoding="utf-8") as f:

        json.dump(
            [c.__dict__ for c in chunks],
            f,
            ensure_ascii=False,
            indent=2
        )

    print("    Paragraphs/tables extracted")
    print(f"    Chunks: {len(chunks)}")
    print(f"    Saved: {output_path}")

    return len(chunks)


# ============================================================
# STANDALONE IMAGES
# ============================================================

def process_standalone_images():

    print("[IMG] Scanning standalone images...")

    image_extensions = {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp"
    }

    images = sorted(
        p
        for p in ROOT.rglob("*")
        if (
            p.is_file()
            and p.suffix.lower() in image_extensions
        )
    )

    image_chunks = []

    for path in images:

        relative_path = str(
            path.relative_to(ROOT)
        )

        visual_id = make_standalone_visual_id(
            relative_path
        )

        output_name = (
            f"{safe_name(path.stem)}"
            f"_{visual_id}"
            f"{path.suffix.lower()}"
        )

        output = VISUALS / output_name

        # Copy original image into visual directory.
        if not output.exists():

            output.write_bytes(
                path.read_bytes()
            )

        image_chunks.append(
            {
                "chunk_id": f"image_{visual_id}",

                "source_file": relative_path,

                "page_number": None,

                "text": (
                    f"Image asset: "
                    f"{relative_path}"
                ),

                "linked_image_id": visual_id
            }
        )

    output_path = (
        PROCESSED /
        "standalone_images_chunks.json"
    )

    with open(output_path, "w", encoding="utf-8") as f:

        json.dump(
            image_chunks,
            f,
            ensure_ascii=False,
            indent=2
        )

    print(
        f"    Standalone images found: "
        f"{len(images)}"
    )

    print(
        f"    Saved: {output_path}"
    )

    return len(image_chunks)


# ============================================================
# MAIN
# ============================================================

def main():

    if not ROOT.exists():

        raise FileNotFoundError(
            f"Archive not found: {ROOT}"
        )

    print("=" * 70)
    print("ASHENLENS - FULL CORPUS INGESTION")
    print("=" * 70)

    print(f"Corpus: {ROOT}")
    print()

    total_chunks = 0

    pdf_stems = set()

    # ========================================================
    # PDFs
    # ========================================================

    pdfs = sorted(
        ROOT.rglob("*.pdf")
    )

    print(
        f"PDF files found: {len(pdfs)}"
    )

    print()

    # PDFs are preferred because
    # they preserve page numbers.

    for path in pdfs:

        pdf_stems.add(
            path.stem.lower()
        )

        total_chunks += process_pdf(
            path
        )

        print()

    # ========================================================
    # MARKDOWN
    # ========================================================

    markdowns = sorted(
        ROOT.rglob("*.md")
    )

    print(
        f"Markdown files found: "
        f"{len(markdowns)}"
    )

    print()

    for path in markdowns:

        total_chunks += process_markdown(
            path
        )

        print()

    # ========================================================
    # TXT
    # ========================================================

    txts = sorted(
        ROOT.rglob("*.txt")
    )

    print(
        f"TXT files found: "
        f"{len(txts)}"
    )

    print()

    for path in txts:

        if path.name.lower() == "readme.txt":
            continue

        total_chunks += process_txt(
            path
        )

        print()

    # ========================================================
    # DOCX
    # ========================================================

    docxs = sorted(
        ROOT.rglob("*.docx")
    )

    docx_only = [
        p
        for p in docxs
        if p.stem.lower()
        not in pdf_stems
    ]

    print(
        f"DOCX files found: "
        f"{len(docxs)}"
    )

    print(
        f"DOCX-only files to process: "
        f"{len(docx_only)}"
    )

    print()

    for path in docx_only:

        total_chunks += process_docx(
            path
        )

        print()

    # ========================================================
    # STANDALONE IMAGES
    # ========================================================

    total_chunks += (
        process_standalone_images()
    )

    print()

    # ========================================================
    # COMPLETE
    # ========================================================

    print("=" * 70)
    print("INGESTION COMPLETE")
    print("=" * 70)

    print(
        f"Total chunks created: "
        f"{total_chunks}"
    )

    print(
        f"Processed data: "
        f"{PROCESSED}"
    )

    print(
        f"Extracted visuals: "
        f"{VISUALS}"
    )

    print("=" * 70)


if __name__ == "__main__":
    main()