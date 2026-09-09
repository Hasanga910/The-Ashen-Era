"""
docx_ingest.py
--------------
Extracts paragraphs, tables, and embedded images from .docx files
(wiki articles / ephemera in the archive that are Word format).

Owner: Data / Extraction lead
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field, asdict
from pathlib import Path

import docx


@dataclass
class ExtractedDocx:
    source_file: str
    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)  # table -> rows -> cells
    image_paths: list[str] = field(default_factory=list)


def extract_docx(docx_path: str | Path, visuals_out_dir: str | Path) -> ExtractedDocx:
    docx_path = Path(docx_path)
    visuals_out_dir = Path(visuals_out_dir)
    visuals_out_dir.mkdir(parents=True, exist_ok=True)

    document = docx.Document(docx_path)
    record = ExtractedDocx(source_file=docx_path.name)

    for para in document.paragraphs:
        if para.text.strip():
            record.paragraphs.append(para.text.strip())

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        record.tables.append(rows)

    # Embedded images live in the docx's media folder inside the zip archive.
    for i, rel in enumerate(document.part.rels.values()):
        if "image" in rel.reltype:
            image_bytes = rel.target_part.blob
            ext = rel.target_ref.split(".")[-1]
            image_path = visuals_out_dir / f"{docx_path.stem}_img{i}.{ext}"
            image_path.write_bytes(image_bytes)
            record.image_paths.append(str(image_path))

    return record


def save_docx_json(record: ExtractedDocx, out_path: str | Path) -> None:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(asdict(record), f, indent=2, ensure_ascii=False)


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python src/ingest/docx_ingest.py <path_to_docx>")
        sys.exit(1)

    result = extract_docx(sys.argv[1], visuals_out_dir="data/visuals")
    print(f"Paragraphs: {len(result.paragraphs)}, Tables: {len(result.tables)}, Images: {len(result.image_paths)}")

    out_json = Path("data/processed") / f"{Path(sys.argv[1]).stem}.json"
    save_docx_json(result, out_json)
    print(f"Saved to {out_json}")
